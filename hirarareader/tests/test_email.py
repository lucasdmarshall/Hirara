"""email_read: .eml, mbox and attachments."""

from __future__ import annotations

import base64
import io
from email.message import EmailMessage

import docx
import pytest

from hirarareader.config import ReaderConfig
from hirarareader.mail import detect, html_to_text
from hirarareader.tools import TOOL_NAMES, Toolset


def _docx_bytes(text: str) -> bytes:
    d = docx.Document()
    d.add_heading("Quarterly report", level=1)
    d.add_paragraph(text)
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def make_email(subject="Invoice 42", html=True) -> bytes:
    msg = EmailMessage()
    msg["From"] = "Ada Lovelace <ada@example.com>"
    msg["To"] = "bob@example.org, Carol <carol@example.org>"
    msg["Subject"] = subject
    msg["Date"] = "Fri, 23 Oct 2026 10:15:00 +0630"
    msg["Message-ID"] = "<abc@example.com>"
    msg["Authentication-Results"] = "mx.example.org; spf=pass; dkim=pass"
    msg.set_content("Hello Bob,\nThe invoice is attached. See https://example.com/pay now.\n")
    if html:
        msg.add_alternative('<p>Hello <b>Bob</b>,</p><p>See <a href="https://evil.test/login">'
                            'your account</a>.</p><script>alert(1)</script>', subtype="html")
    msg.add_attachment(_docx_bytes("Revenue grew 12%."), maintype="application",
                       subtype="vnd.openxmlformats-officedocument.wordprocessingml.document",
                       filename="report.docx")
    msg.add_attachment(b"a,b\n1,2\n", maintype="text", subtype="csv", filename="data.csv")
    msg.add_attachment(b"\x89PNG....", maintype="image", subtype="png", filename="logo.png")
    return bytes(msg)


@pytest.fixture
def ts():
    return Toolset(ReaderConfig())


async def test_eml_headers_body_links_and_attachments(ts, tmp_path):
    p = tmp_path / "m.eml"
    p.write_bytes(make_email())
    r = await ts.email_read(file_path=str(p))
    assert r["error"] is None and r["format"] == "eml" and r["count"] == 1
    m = r["messages"][0]
    assert m["subject"] == "Invoice 42"
    assert m["from"] == [{"name": "Ada Lovelace", "email": "ada@example.com"}]
    assert [a["email"] for a in m["to"]] == ["bob@example.org", "carol@example.org"]
    assert m["date"] == "2026-10-23T10:15:00+06:30"
    assert m["body_text"].startswith("Hello Bob,")
    assert set(m["links"]) == {"https://evil.test/login", "https://example.com/pay"}
    assert m["authentication_results"] == ["mx.example.org; spf=pass; dkim=pass"]
    names = {a["filename"]: a for a in m["attachments"]}
    assert "Revenue grew 12%" in names["report.docx"]["text"]
    assert names["data.csv"]["text"] == "a,b\n1,2\n"
    assert names["logo.png"]["text"] is None and len(names["logo.png"]["sha256"]) == 64
    assert "Ada" not in r["text"] and "ada@example.com" in r["text"]


async def test_html_only_body_is_converted(ts):
    msg = EmailMessage()
    msg["From"] = "a@b.c"
    msg["Subject"] = "html"
    msg.set_content("<h1>Title</h1><p>Line one</p><style>x{}</style>", subtype="html")
    r = await ts.email_read(file_base64=base64.b64encode(bytes(msg)).decode(), filename="x.eml",
                            include_html=True)
    m = r["messages"][0]
    assert m["body_text"] == "Title\n\nLine one"
    assert m["body_html"].startswith("<h1>")


async def test_mbox_lists_messages(ts):
    box = b"".join(b"From someone@example.com Fri Oct 23 10:00:00 2026\n" + make_email(f"Msg {i}", html=False)
                   + b"\n" for i in range(3))
    r = await ts.email_read(file_base64=base64.b64encode(box).decode(), max_messages=2)
    assert r["format"] == "mbox" and r["count"] == 2 and r["total"] == 3
    assert [m["subject"] for m in r["messages"]] == ["Msg 0", "Msg 1"]
    assert "1 more message" in r["text"]


async def test_skip_reading_attachments(ts):
    r = await ts.email_read(file_base64=base64.b64encode(make_email()).decode(),
                            read_attachments=False)
    assert all(a["text"] is None for a in r["messages"][0]["attachments"])


async def test_errors_and_detection(ts):
    assert "provide exactly one" in (await ts.email_read())["error"]
    r = await ts.email_read(file_base64=base64.b64encode(b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1junk").decode())
    assert r["error"]  # detected as .msg, then rejected as unreadable
    assert detect(b"From x\n", None) == "mbox" and detect(b"Subject: x\n\nhi", "a.eml") == "eml"
    r = await ts.email_read(file_base64=base64.b64encode(b"just some words").decode())
    assert "not an email" in r["error"]


def test_html_to_text_drops_scripts():
    text, links = html_to_text('<p>a</p><script>bad()</script><a href="https://x.y">x</a>')
    assert "bad" not in text and links == ["https://x.y"]


def test_tool_names():
    assert TOOL_NAMES == ("office_read", "email_read")
